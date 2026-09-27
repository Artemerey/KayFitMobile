import json
import shlex
import subprocess
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUT = Path(__file__).resolve().parents[1] / "specs" / "KayFit_отзывы_и_ошибки_2026-09-25.xlsx"
SSH = [
    "ssh", "-i", "/Users/user/.ssh/id_kayfit_deploy", "-p", "2222",
    "-o", "BatchMode=yes", "root@89.23.99.129",
]


def query(sql):
    remote = [
        "docker", "exec", "postgres", "psql", "-U", "postgres", "-d", "calories",
        "-At", "-c", f"SELECT COALESCE(json_agg(x),'[]'::json) FROM ({sql}) x",
    ]
    raw = subprocess.run(SSH + [shlex.join(remote)], check=True, capture_output=True, text=True).stdout
    return json.loads(raw)


feedback = query("""
SELECT row_number() OVER (ORDER BY f.created_at,f.id) AS case_no,
       to_char(f.created_at AT TIME ZONE 'Europe/Moscow','YYYY-MM-DD HH24:MI:SS') AS time_msk,
       f.target_type, f.rating, array_to_string(f.reason_codes, ', ') AS reasons,
       COALESCE(f.comment,'') AS comment, f.source, f.context,
       COALESCE(f.app_version,'') AS app_version,
       COALESCE(f.platform,'') AS platform, COALESCE(f.locale,'') AS locale,
       COALESCE(f.traceability_status,'') AS traceability,
       CASE WHEN f.revision_of IS NULL THEN 'нет' ELSE 'да' END AS is_revision,
       CASE WHEN f.comment ILIKE '%QA%' OR f.comment ILIKE '%Synthetic production E2E%'
            THEN 'тестовый' ELSE 'пользовательский/не помечен' END AS data_kind
FROM user_feedback f
ORDER BY f.created_at,f.id
""")

recognitions = query("""
SELECT row_number() OVER (ORDER BY f.created_at,f.id,i.ordinal) AS row_no,
       (SELECT count(*) FROM user_feedback f2
        WHERE (f2.created_at,f2.id) <= (f.created_at,f.id)) AS case_no,
       to_char(f.created_at AT TIME ZONE 'Europe/Moscow','YYYY-MM-DD HH24:MI:SS') AS time_msk,
       f.rating, array_to_string(f.reason_codes, ', ') AS reasons, f.source,
       i.ordinal + 1 AS item_no,
       COALESCE(i.recognized_snapshot->>'name','') AS food,
       NULLIF(i.recognized_snapshot->>'weight_grams','')::numeric AS weight_g,
       NULLIF(i.recognized_snapshot->>'calories','')::numeric AS calories,
       NULLIF(i.recognized_snapshot->>'protein','')::numeric AS protein,
       NULLIF(i.recognized_snapshot->>'fat','')::numeric AS fat,
       NULLIF(i.recognized_snapshot->>'carbs','')::numeric AS carbs,
       COALESCE(i.recognized_snapshot->>'meal_type','') AS meal_type,
       COALESCE(i.recognized_snapshot->>'source','') AS nutrition_source,
       COALESCE(f.comment,'') AS user_comment,
       COALESCE(f.context::text,'{}') AS feedback_context
FROM user_feedback f
JOIN meal_processing_item i ON i.run_id=f.meal_run_id
WHERE f.target_type='meal_save'
ORDER BY f.created_at,f.id,i.ordinal
""")

logs = [
    ["2026-09-23 17:49:38", "POST /api/meals/copy-batch", "200 OK", "Сервер принял запрос на копирование"],
    ["2026-09-23 17:50:05", "DELETE /api/meals/5427", "500", "meal processing evidence cannot be rewritten"],
    ["2026-09-23 17:50:21", "POST /api/meals/copy-batch", "200 OK", "Сервер принял ещё один запрос на копирование"],
    ["2026-09-23 17:50:51", "DELETE /api/meals/5471", "500", "meal processing evidence cannot be rewritten"],
    ["2026-09-23 17:50:53", "DELETE /api/meals/5427", "500", "meal processing evidence cannot be rewritten"],
    ["2026-09-23 17:51:16", "DELETE /api/meals/5471", "500", "meal processing evidence cannot be rewritten"],
    ["2026-09-23 17:51:31", "DELETE /api/meals/5472", "500", "meal processing evidence cannot be rewritten"],
    ["2026-09-23 17:51:36", "DELETE /api/meals/5471", "500", "meal processing evidence cannot be rewritten"],
    ["2026-09-23 17:53:13", "Отзыв пользователя", "dislike / other", "copying to original date, not new selected date, then can’t delete wrong copy"],
]

plans = [
    [1, "Нельзя удалить некоторые блюда", "Сервер пытается стереть блюдо, но защищённая история запрещает это. Получается ошибка 500.", "Не стирать строку. Помечать блюдо удалённым и скрывать его от пользователя.", "Срочно"],
    [2, "Копия может появиться не в ожидаемом дне", "Сервер сохранил дату, которую прислал Flutter. Нужно проверить календарь и передачу даты.", "Вернуть из API фактическую дату и открыть во Flutter именно этот день.", "Срочно"],
    [3, "Неправильные калории и БЖУ", "Пользователь ставит минус, но правильные цифры система не получает.", "После минуса предложить указать правильный вес, калории или БЖУ. Из исправлений сделать тесты.", "Высоко"],
    [4, "Неправильно названо блюдо", "Есть жалобы на название, но часто нет исходного текста или фото для проверки.", "Сохранять безопасный исходный контекст и спрашивать уточнение, если система не уверена.", "Высоко"],
    [5, "Распознавание долгое", "Есть 2 жалобы, но настоящее время распознавания не записано.", "Измерять время от отправки до результата и прикладывать его к отзыву.", "Высоко"],
    [6, "План даёт слишком много калорий", "7 из 10 плохих оценок плана содержат эту причину.", "Проверить формулу, цель, активность, единицы и ограничения калорий.", "Высоко"],
    [7, "Мало данных об отзыве", "Нет версии приложения, платформы и режима распознавания.", "Добавлять эти поля автоматически без личных данных.", "Высоко"],
]


wb = Workbook()
ws = wb.active
ws.title = "Главное"
ws.append(["KayFit — что нашли в отзывах"])
ws.append([])
ws.append(["Проверено отзывов", 75])
ws.append(["Положительных", 47])
ws.append(["Отрицательных", 28])
ws.append(["Точная фраза пользователя", "copying to original date, not new selected date, then can’t delete wrong copy"])
ws.append(["Перевод", "Копируется на исходную дату, а не на новую выбранную; потом неправильную копию нельзя удалить."])
ws.append(["Где найдено", "Production user_feedback, 2026-09-23 17:53:13 MSK, кейс №75"])
ws.append(["Что подтверждают серверные логи", "Пользователь несколько раз пытался удалить блюда. DELETE вернулся с ошибкой 500."])
ws.append(["Главная причина", "Обычное удаление конфликтует с защищённой историей распознавания."])
ws.append(["Простое решение", "Помечать блюдо удалённым и скрывать его, но не стирать техническую историю."])

ws = wb.create_sheet("Все отзывы")
headers = ["№", "Время MSK", "Объект", "Оценка", "Причины", "Комментарий пользователя", "Источник", "Контекст", "Версия", "Платформа", "Язык", "Трассировка", "Повторная оценка", "Тип данных"]
ws.append(headers)
for r in feedback:
    values = [r[k] for k in ["case_no", "time_msk", "target_type", "rating", "reasons", "comment", "source", "context", "app_version", "platform", "locale", "traceability", "is_revision", "data_kind"]]
    values[7] = json.dumps(values[7], ensure_ascii=False)
    ws.append(values)

ws = wb.create_sheet("Примеры распознавания")
headers = ["Строка", "№ отзыва", "Время MSK", "Оценка", "Причины", "Источник", "Позиция", "Что распознано", "Вес, г", "Ккал", "Белки", "Жиры", "Углеводы", "Приём пищи", "Источник данных", "Комментарий", "Контекст отзыва"]
ws.append(headers)
for r in recognitions:
    ws.append([r[k] for k in ["row_no", "case_no", "time_msk", "rating", "reasons", "source", "item_no", "food", "weight_g", "calories", "protein", "fat", "carbs", "meal_type", "nutrition_source", "user_comment", "feedback_context"]])

ws = wb.create_sheet("Подтверждающие логи")
ws.append(["Время MSK", "Событие", "Результат", "Что это значит"])
for row in logs:
    ws.append(row)

ws = wb.create_sheet("Что исправить")
ws.append(["№", "Проблема", "Почему это происходит", "Простое решение", "Приоритет"])
for row in plans:
    ws.append(row)

ws = wb.create_sheet("Ограничения")
ws.append(["Что важно понимать"])
ws.append(["Для старых отзывов система не сохраняла связанный снимок распознавания. Поэтому на листе «Примеры распознавания» есть только реально доступные примеры; отсутствующие данные не выдуманы."])
ws.append(["Отзыв показывает мнение пользователя, но не всегда сообщает правильное блюдо или правильные цифры. Такие случаи нужно превратить в исправления от пользователя и тесты."])
ws.append(["Тестовые QA/E2E отзывы помечены на листе «Все отзывы» и не должны смешиваться с обычной продуктовой статистикой."])

navy = "1F4E78"
blue = "D9EAF7"
red = "F4CCCC"
green = "D9EAD3"
yellow = "FFF2CC"
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(row=r, column=col).value or "") for r in range(1, min(ws.max_row, 100) + 1)]
        width = min(55, max(10, max(map(len, values)) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width

for row in wb["Все отзывы"].iter_rows(min_row=2):
    if row[3].value == "dislike":
        for c in row:
            c.fill = PatternFill("solid", fgColor=red)
    elif row[3].value == "like":
        for c in row:
            c.fill = PatternFill("solid", fgColor=green)
    if row[0].value == 75:
        for c in row:
            c.fill = PatternFill("solid", fgColor=yellow)
            c.font = Font(bold=True)

for row in wb["Примеры распознавания"].iter_rows(min_row=2):
    if row[3].value == "dislike":
        for c in row:
            c.fill = PatternFill("solid", fgColor=red)

wb["Главное"].column_dimensions["A"].width = 34
wb["Главное"].column_dimensions["B"].width = 100
wb["Ограничения"].column_dimensions["A"].width = 120
wb.save(OUT)

check = load_workbook(OUT, data_only=False)
assert len(feedback) == 75
assert check["Все отзывы"].max_row == 76
assert check["Примеры распознавания"].max_row == len(recognitions) + 1
assert not any(
    isinstance(c.value, str) and c.value.startswith("#")
    for s in check.worksheets for row in s.iter_rows() for c in row
)
print(json.dumps({"file": str(OUT), "feedback": len(feedback), "recognition_rows": len(recognitions), "sheets": check.sheetnames}, ensure_ascii=False))
